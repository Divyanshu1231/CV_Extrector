"""Upload formats (PDF / image / Word) and export formats (PDF / Word / image)."""
import io
import json
import shutil
import zipfile

import fitz
import pytest
from fastapi.testclient import TestClient

from app.main import app
from tests.test_robust import title_case_cv, needs_tesseract

client = TestClient(app)
needs_office = pytest.mark.skipif(shutil.which("soffice") is None and shutil.which("libreoffice") is None,
                                  reason="LibreOffice not installed")


def _jpeg_of_cv(dpi=130) -> bytes:
    return fitz.open(stream=title_case_cv(), filetype="pdf")[0].get_pixmap(dpi=dpi).tobytes("jpeg")


def _text(pdf: bytes) -> str:
    return "\n".join(p.get_text() for p in fitz.open(stream=pdf, filetype="pdf"))


@needs_tesseract
def test_single_image_upload_is_processed():
    r = client.post("/api/process-cv", files={"file": ("cv photo.jpg", _jpeg_of_cv(), "image/jpeg")})
    assert r.status_code == 200, r.text
    assert r.headers["x-cv-code"] == "DR. AKS 3210"
    rep = json.loads(r.headers["x-report"])
    assert rep["source"] == "image" and rep["verified"] is True


@needs_tesseract
def test_two_images_become_two_pages():
    imgs = [("files", ("p1.jpg", _jpeg_of_cv(), "image/jpeg")), ("files", ("p2.jpg", _jpeg_of_cv(), "image/jpeg"))]
    r = client.post("/api/process-cv", files=imgs)
    assert r.status_code == 200, r.text
    assert fitz.open(stream=r.content, filetype="pdf").page_count == 2


@needs_office
def test_docx_upload_is_converted_and_protected():
    import docx
    d = docx.Document()
    d.add_paragraph("Dr. Neha Gupta").runs[0].font.size = docx.shared.Pt(26)
    d.add_paragraph("Mobile: +91 98123 45678 | neha.gupta@example.com")
    d.add_paragraph("Consultant, City Hospital, 2019 - Present")
    buf = io.BytesIO(); d.save(buf)
    r = client.post("/api/process-cv", files={"file": ("Neha_CV.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert r.headers["x-cv-code"] == "DR. NG 5678"
    t = _text(r.content)
    assert "Neha" not in t and "98123" not in t and "neha.gupta" not in t
    assert "City Hospital" in t


def test_unsupported_and_mixed_uploads_are_rejected_clearly():
    r = client.post("/api/process-cv", files={"file": ("x.exe", b"MZ12345678", "application/octet-stream")})
    assert r.status_code == 400 and r.json()["detail"]["code"] == "UNSUPPORTED"
    r = client.post("/api/process-cv", files=[("files", ("a.pdf", title_case_cv(), "application/pdf")),
                                               ("files", ("b.jpg", _jpeg_of_cv(), "image/jpeg"))])
    assert r.status_code == 400 and r.json()["detail"]["code"] == "MIXED_FILES"


def test_export_word_is_editable_and_image_export_works():
    pdf = client.post("/api/process-cv", files={"file": ("a.pdf", title_case_cv(), "application/pdf")}).content
    r = client.post("/api/export", files={"file": ("a.pdf", pdf, "application/pdf")},
                    data={"format": "docx", "base_name": "DR_AKS_3210", "cv_code": "DR. AKS 3210"})
    assert r.status_code == 200 and r.headers["x-editable"] == "1"
    import docx
    doc = docx.Document(io.BytesIO(r.content))
    body = " ".join(p.text for p in doc.paragraphs)
    assert "Apollo Hospital" in body and "Sharma" not in body
    assert "DR. AKS 3210" in doc.sections[0].footer.paragraphs[0].text

    r = client.post("/api/export", files={"file": ("a.pdf", pdf, "application/pdf")}, data={"format": "png", "base_name": "X"})
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content[:4] == b"\x89PNG"

    two = fitz.open(); two.insert_pdf(fitz.open(stream=pdf, filetype="pdf")); two.insert_pdf(fitz.open(stream=pdf, filetype="pdf"))
    r = client.post("/api/export", files={"file": ("a.pdf", two.tobytes(), "application/pdf")}, data={"format": "jpg", "base_name": "X"})
    assert r.headers["content-type"] == "application/zip"
    assert len(zipfile.ZipFile(io.BytesIO(r.content)).namelist()) == 2
