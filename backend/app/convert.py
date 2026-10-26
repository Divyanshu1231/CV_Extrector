"""Input/output format support.

Input  : PDF, images (JPG/PNG/WEBP/BMP/TIFF/GIF - several photos = several pages),
         Word (DOCX, DOC, also ODT/RTF through LibreOffice).
Output : PDF (primary), Word (editable), PNG/JPG images.

Everything is normalised to a PDF first, so the privacy pipeline in
processor.py stays identical for every kind of upload.
"""
from __future__ import annotations

import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile

import fitz

from .processor import WATERMARK, DetectionError, page_needs_ocr

logger = logging.getLogger("cv-privacy")

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
WORD_EXTS = {".docx", ".doc", ".odt", ".rtf"}
ALLOWED_EXTS = {".pdf"} | IMAGE_EXTS | WORD_EXTS
MAX_IMAGE_SIDE = 3200       # px; bigger phone photos are scaled down (OCR does not need more)
OFFICE_TIMEOUT_S = 120


# --------------------------------------------------------------------------- detection
def ext_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def detect_kind(filename: str, data: bytes) -> str:
    """Return 'pdf' | 'image' | 'word' by content (magic bytes), falling back to the extension."""
    head = data[:16]
    if data[:1024].lstrip().startswith(b"%PDF"):
        return "pdf"
    if head.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"BM")) or head[:4] in (b"II*\x00", b"MM\x00*") \
            or (head[:4] == b"RIFF" and data[8:12] == b"WEBP"):
        return "image"
    if head.startswith(b"PK"):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
            if any(n.startswith("word/") for n in names):
                return "word"
            if "mimetype" in names:
                return "word"  # ODT
        except Exception:
            pass
    if head.startswith(b"\xd0\xcf\x11\xe0") or data[:5] == b"{\\rtf":
        return "word"
    ext = ext_of(filename)
    if ext in IMAGE_EXTS:
        return "image"
    if ext in WORD_EXTS:
        return "word"
    return "unknown"


def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def office_available() -> bool:
    return soffice_path() is not None


def docx_export_available() -> bool:
    try:
        import pdf2docx  # noqa: F401
        return True
    except Exception:
        return False


# --------------------------------------------------------------------------- inputs -> PDF
def images_to_pdf(images: list[bytes]) -> bytes:
    """One page per image, page size follows the image aspect ratio (A4 width)."""
    from PIL import Image, ImageOps

    if not images:
        raise DetectionError("EMPTY_FILE", "No image was uploaded.")
    doc = fitz.open()
    for raw in images:
        try:
            img = Image.open(io.BytesIO(raw))
            img = ImageOps.exif_transpose(img)           # phone photos: honour rotation
            if getattr(img, "n_frames", 1) > 1:
                img.seek(0)
            if img.mode in ("RGBA", "LA", "P"):
                bg = Image.new("RGB", img.size, (255, 255, 255))
                rgba = img.convert("RGBA")
                bg.paste(rgba, mask=rgba.split()[-1])
                img = bg
            else:
                img = img.convert("RGB")
        except Exception:
            raise DetectionError("BAD_IMAGE", "One of the images could not be opened. Please upload a clear JPG or PNG.")
        w, h = img.size
        scale = min(1.0, MAX_IMAGE_SIDE / max(w, h))
        if scale < 1.0:
            img = img.resize((int(w * scale), int(h * scale)), Image.LANCZOS)
            w, h = img.size
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=90, optimize=True)
        page_w = 595.0
        page = doc.new_page(width=page_w, height=page_w * h / w)
        page.insert_image(page.rect, stream=buf.getvalue())
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def word_to_pdf(data: bytes, filename: str) -> bytes:
    exe = soffice_path()
    if exe is None:
        raise DetectionError(
            "OFFICE_UNAVAILABLE",
            "Word files cannot be converted on this server yet (LibreOffice is not installed). "
            "Please upload the CV as PDF or image instead.",
        )
    ext = ext_of(filename) if ext_of(filename) in WORD_EXTS else ".docx"
    with tempfile.TemporaryDirectory(prefix="cvword_") as tmp:
        src = os.path.join(tmp, "cv" + ext)
        with open(src, "wb") as fh:
            fh.write(data)
        profile = os.path.join(tmp, "profile")
        cmd = [
            exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
            "--nolockcheck", "--convert-to", "pdf:writer_pdf_Export", "--outdir", tmp, src,
        ]
        try:
            subprocess.run(cmd, check=True, timeout=OFFICE_TIMEOUT_S, capture_output=True,
                           env={**os.environ, "HOME": tmp})
        except subprocess.TimeoutExpired:
            raise DetectionError("WORD_TIMEOUT", "The Word file took too long to convert. Please upload it as PDF.")
        except subprocess.CalledProcessError as exc:
            logger.warning("soffice failed: %s", (exc.stderr or b"")[:300])
            raise DetectionError("WORD_CONVERT_FAILED", "This Word file could not be converted. Please save it as PDF and upload again.")
        out = os.path.join(tmp, "cv.pdf")
        if not os.path.exists(out):
            raise DetectionError("WORD_CONVERT_FAILED", "This Word file could not be converted. Please save it as PDF and upload again.")
        with open(out, "rb") as fh:
            return fh.read()


# --------------------------------------------------------------------------- PDF -> outputs
def _page_has_text(page: fitz.Page) -> bool:
    """True for pages with a real text layer (typed CVs), False for scans/pictures."""
    return not page_needs_ocr(page)


def pdf_to_images(pdf_bytes: bytes, fmt: str = "png", dpi: int = 170, base: str = "CV") -> tuple[bytes, str, str]:
    """Return (bytes, filename, mime). One page -> a single image, several pages -> ZIP of images."""
    fmt = "jpg" if fmt.lower() in ("jpg", "jpeg") else "png"
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        files = []
        for i, page in enumerate(doc, start=1):
            pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csRGB, alpha=False)
            files.append((f"{base}_page{i}.{fmt}", pix.tobytes("jpeg" if fmt == "jpg" else "png")))
        mime = "image/jpeg" if fmt == "jpg" else "image/png"
        if len(files) == 1:
            return files[0][1], f"{base}.{fmt}", mime
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
            for name, content in files:
                z.writestr(name, content)
        return buf.getvalue(), f"{base}_images.zip", "application/zip"
    finally:
        doc.close()


def _images_docx(pdf_bytes: bytes) -> bytes:
    """Fallback Word file: each page as a full-width picture (not editable text)."""
    from docx import Document
    from docx.shared import Mm

    src = fitz.open(stream=pdf_bytes, filetype="pdf")
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(8)
    for i, page in enumerate(src):
        pix = page.get_pixmap(dpi=150, colorspace=fitz.csRGB, alpha=False)
        d.add_picture(io.BytesIO(pix.tobytes("png")), width=Mm(194))
        if i < src.page_count - 1:
            d.add_page_break()
    src.close()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _stamp_docx(docx_bytes: bytes, cv_code: str) -> bytes:
    """PDF->Word drops the diagonal watermark, so the CV Code and agency line are
    written into the footer of every section instead."""
    if not cv_code:
        return docx_bytes
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor

        d = Document(io.BytesIO(docx_bytes))
        for sec in d.sections:
            footer = sec.footer
            footer.is_linked_to_previous = False
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.text = ""
            run = para.add_run(f"CV Code: {cv_code}   |   {WATERMARK}")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        out = io.BytesIO()
        d.save(out)
        return out.getvalue()
    except Exception as exc:
        logger.warning("Could not stamp docx footer: %s", exc)
        return docx_bytes


def pdf_to_docx(pdf_bytes: bytes, cv_code: str = "") -> tuple[bytes, bool]:
    """Return (docx_bytes, editable). Text CVs become an editable Word document;
    scanned/image CVs (no text to edit) fall back to pictures of the pages."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    try:
        editable = any(_page_has_text(p) for p in doc)
    finally:
        doc.close()
    if editable and docx_export_available():
        try:
            from pdf2docx import Converter

            with tempfile.TemporaryDirectory(prefix="cvdocx_") as tmp:
                src = os.path.join(tmp, "in.pdf")
                dst = os.path.join(tmp, "out.docx")
                with open(src, "wb") as fh:
                    fh.write(pdf_bytes)
                cv = Converter(src)
                try:
                    cv.convert(dst, start=0, end=None)
                finally:
                    cv.close()
                with open(dst, "rb") as fh:
                    return _stamp_docx(fh.read(), cv_code), True
        except Exception as exc:
            logger.warning("pdf2docx failed, using picture fallback: %s", exc)
    return _stamp_docx(_images_docx(pdf_bytes), cv_code), False
